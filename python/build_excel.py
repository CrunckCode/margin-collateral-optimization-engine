"""Build the Excel version: live formulas for IM and coverage; the LP solution is pasted from Python."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from margin_engine import config as C
from margin_engine.collateral import build_inventory
from margin_engine.optimizer import optimize
from margin_engine.simm_style import headline_im

OUT = C.PROJECT_DIR / "excel" / "Margin_Collateral_Optimization.xlsx"
BOLD = Font(bold=True)


def put(ws, row, values, bold=False):
    for j, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=j, value=v)
        if bold:
            cell.font = BOLD


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def main():
    im = headline_im()
    names = list(im)
    inv = build_inventory()
    alloc = optimize(inv, im[names[0]]["im"], im[names[1]]["im"])
    wb = Workbook()

    ws = wb.active
    ws.title = "README"
    lines = ["Margin and Collateral Optimization Engine (Excel companion)", "",
             "Inputs: all assumptions. IM: SIMM-style initial margin, live formulas. Collateral: inventory and haircuts.",
             "Allocation: the linear-program solution pasted from Python (Solver is not required), with live checks.",
             "Scenarios: static results from the Python scenario runs.",
             "SIMM-style means simplified and illustrative. It is not ISDA SIMM and not ISDA certified.",
             "To re-optimize after changing inputs, run python/run_all.py and paste the new allocation, or use Excel Solver."]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = BOLD if i == 1 else Font()
    widths(ws, {"A": 110})

    wi = wb.create_sheet("Inputs")
    put(wi, 1, ["Market snapshot", "Value"], True)
    rows = [("AAPL price", C.PRICES["AAPL"]), ("JPM price", C.PRICES["JPM"]), ("XOM price", C.PRICES["XOM"]),
            ("EURUSD", C.EURUSD), ("US 10Y yield", C.US10Y_YIELD), ("Duration proxy (years)", C.RATE_DURATION_YEARS)]
    for i, r in enumerate(rows, 2):
        put(wi, i, r)
    put(wi, 9, ["Netting set positions", names[0], names[1]], True)
    sa, sb = C.NETTING_SETS[names[0]], C.NETTING_SETS[names[1]]
    pos = [("AAPL shares", sa["shares"].get("AAPL", 0), sb["shares"].get("AAPL", 0)),
           ("JPM shares", sa["shares"].get("JPM", 0), sb["shares"].get("JPM", 0)),
           ("XOM shares", sa["shares"].get("XOM", 0), sb["shares"].get("XOM", 0)),
           ("EUR notional", sa["eur_notional"], sb["eur_notional"]),
           ("USD rate notional", sa["rate_notional"], sb["rate_notional"])]
    for i, r in enumerate(pos, 10):
        put(wi, i, r)
    put(wi, 16, ["Risk weights", "Value"], True)
    put(wi, 17, ["Equity", C.RISK_WEIGHTS["equity"]])
    put(wi, 18, ["FX", C.RISK_WEIGHTS["fx"]])
    put(wi, 19, ["Rate", C.RISK_WEIGHTS["rate"]])
    put(wi, 20, ["Cross-class correlation", C.CROSS_CLASS_CORRELATION])
    put(wi, 21, ["Cash opportunity cost (bps)", C.CASH_OPPORTUNITY_COST_BPS])
    widths(wi, {"A": 30, "B": 32, "C": 32})

    wm = wb.create_sheet("IM")
    put(wm, 1, ["Item", names[0], names[1]], True)
    put(wm, 2, ["Equity delta (USD)", "=Inputs!B10*Inputs!$B$2+Inputs!B11*Inputs!$B$3+Inputs!B12*Inputs!$B$4",
                "=Inputs!C10*Inputs!$B$2+Inputs!C11*Inputs!$B$3+Inputs!C12*Inputs!$B$4"])
    put(wm, 3, ["FX delta (USD)", "=Inputs!B13*Inputs!$B$5", "=Inputs!C13*Inputs!$B$5"])
    put(wm, 4, ["Rate exposure (USD, duration-style)", "=Inputs!B14*Inputs!$B$7/(1+Inputs!$B$6)", "=Inputs!C14*Inputs!$B$7/(1+Inputs!$B$6)"])
    put(wm, 5, ["Weighted equity", "=ABS(B2)*Inputs!$B$17", "=ABS(C2)*Inputs!$B$17"])
    put(wm, 6, ["Weighted FX", "=ABS(B3)*Inputs!$B$18", "=ABS(C3)*Inputs!$B$18"])
    put(wm, 7, ["Weighted rate", "=ABS(B4)*Inputs!$B$19", "=ABS(C4)*Inputs!$B$19"])
    put(wm, 8, ["Initial margin", "=SQRT(B5^2+B6^2+B7^2+2*Inputs!$B$20*(B5*B6+B5*B7+B6*B7))",
                "=SQRT(C5^2+C6^2+C7^2+2*Inputs!$B$20*(C5*C6+C5*C7+C6*C7))"], True)
    put(wm, 10, ["Total initial margin", "=B8+C8"], True)
    widths(wm, {"A": 40, "B": 32, "C": 32})

    wc = wb.create_sheet("Collateral")
    put(wc, 1, ["Asset", "Market value", "Haircut", "Opportunity cost (bps)", "Eligible A", "Eligible B",
                "Post-haircut value", "Cost per unit of coverage (bps)"], True)
    for i, r in enumerate(inv.itertuples(index=False), 2):
        put(wc, i, [r.asset, r.market_value, r.haircut, r.opp_cost_bps, "Yes" if r.elig_A else "No",
                    "Yes" if r.elig_B else "No", f"=B{i}*(1-C{i})", f"=D{i}/(1-C{i})"])
    widths(wc, {"A": 24, "B": 18, "C": 12, "D": 22, "E": 12, "F": 12, "G": 20, "H": 30})

    wa = wb.create_sheet("Allocation")
    put(wa, 1, ["Asset", "Posted to A (market value)", "Posted to B (market value)", "Total posted", "Within inventory?", "Annual cost"], True)
    n = len(alloc.table)
    for i, r in enumerate(alloc.table.itertuples(index=False), 2):
        put(wa, i, [r.asset, float(r.to_A), float(r.to_B), f"=B{i}+C{i}",
                    f'=IF(D{i}<=Collateral!B{i}+0.01,"OK","BREACH")', f"=D{i}*Collateral!D{i}/10000"])
    last = n + 1
    put(wa, last + 2, ["Coverage to A (post-haircut)", f"=SUMPRODUCT(B2:B{last},1-Collateral!C2:C{last})"], True)
    put(wa, last + 3, ["Initial margin A", "=IM!B8"])
    put(wa, last + 4, ["A covered?", f'=IF(B{last+2}>=B{last+3}-1,"OK","SHORT")'])
    put(wa, last + 5, ["Coverage to B (post-haircut)", f"=SUMPRODUCT(C2:C{last},1-Collateral!C2:C{last})"], True)
    put(wa, last + 6, ["Initial margin B", "=IM!C8"])
    put(wa, last + 7, ["B covered?", f'=IF(B{last+5}>=B{last+6}-1,"OK","SHORT")'])
    put(wa, last + 9, ["Optimized annual cost", f"=SUM(F2:F{last})"], True)
    put(wa, last + 10, ["Cash-only annual cost", "=IM!B10*Inputs!B21/10000"], True)
    put(wa, last + 11, ["Annual saving", f"=B{last+10}-B{last+9}"], True)
    put(wa, last + 12, ["Saving percent", f"=B{last+11}/B{last+10}"], True)
    widths(wa, {"A": 34, "B": 28, "C": 28, "D": 20, "E": 20, "F": 18})

    ws2 = wb.create_sheet("Scenarios")
    r = 1
    for title, csv in [("Funding rate sweep", "scenario_funding_sweep.csv"), ("Haircut stress", "scenario_haircut_stress.csv"),
                       ("Initial margin scaling", "scenario_im_scaling.csv"), ("Treasury scarcity", "scenario_treasury_scarcity.csv"),
                       ("Eligibility cases", "scenario_eligibility.csv")]:
        df = pd.read_csv(C.OUTPUT_DIR / csv)
        put(ws2, r, [title + " (static, from Python)"], True)
        put(ws2, r + 1, list(df.columns), True)
        for k, row in enumerate(df.itertuples(index=False), r + 2):
            put(ws2, k, [None if (not isinstance(v, str) and pd.isna(v)) else (bool(v) if str(v) in ("True", "False") else v) for v in row])
        r += len(df) + 4
    widths(ws2, {"A": 42, "B": 22, "C": 22, "D": 22, "E": 22, "F": 22, "G": 22})

    wb.properties.creator = "Deepak Chaudhary"
    wb.properties.lastModifiedBy = "Deepak Chaudhary"
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
